"""`_force_iterative_calc` must work on the files it was WRITTEN for.

WHY THIS EXISTS. The function forces Excel's iterative-calculation flag on before
the scorer's LibreOffice recalc, because several confirmed SSB v2 false negatives
are circular-reference "scenario toggle" cells that LibreOffice cannot resolve at
all without it — it returns #VALUE! regardless of what the solver did.

Its own docstring records the motivating measurement: of 10 sibling bug-variant
inputs sharing one base template, **8 lack the flag entirely**. openpyxl exposes a
missing `<calcPr>` element as `wb.calculation is None`, so the original
implementation's very first statement — `if wb.calculation.iterate` — raised
AttributeError on exactly those 8. The `except Exception` around it logged a
warning and returned, so the fix silently did nothing for the majority of its
target population, and the warning's wording ("could not open/save") read like an
I/O hiccup rather than a lost correctness guarantee.

Found 2026-08-02 from a live rollout log, then reproduced directly against the
pristine `Financial_Model/02_01` dataset input. These tests pin all three branches
so the None case cannot regress back into a silent no-op.
"""

from __future__ import annotations

import logging
import os
import shutil
import sys
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402
from openpyxl.workbook.properties import CalcProperties  # noqa: E402

from src.agent_profiles.ssbv2_agent.scorer import _force_iterative_calc  # noqa: E402

# A REAL dataset input, not a synthesised one. The bug is a property of how this
# dataset's files were produced (writer "Aptura", no calcPr element), so a
# hand-built fixture could pass while the real population still failed.
#
# RESOLVED FROM $SSBV2_ROOT, NOT FROM A SIBLING OF THE CLONE. This was
# `ROOT.parent / "SpreadsheetBench-v2/data/spreadsheetbench-v2/..."` -- the directory
# CONTAINING the repo, which happens to hold the dataset on the machine these tests were
# written on and nowhere else. Measured: in-tree `pytest -q` -> 425 passed; the same
# suite against a `git archive HEAD` copy -> 423 passed, 2 SKIPPED "SpreadsheetBench-v2
# dataset not present". So on every fresh clone the only two guards against this fix
# regressing into a silent no-op silently did not run, and the suite was green either way.
# resolve_base() is used because two checkout layouts are in the wild (<root>/data/... and
# <root>/...), which is the same reason the runtime code calls it.
_SSBV2_ROOT = os.environ.get("SSBV2_ROOT", "").strip()
_REL = "Financial_Model/spreadsheet/02_Project PP/02_01_PP_input.xlsx"

if _SSBV2_ROOT:
    from src.agent_profiles.ssbv2_agent.dataset import resolve_base  # noqa: E402
    NO_CALCPR_INPUT = Path(resolve_base(_SSBV2_ROOT)) / _REL
else:
    NO_CALCPR_INPUT = Path("/nonexistent")

# SKIP only when the dataset was never configured; FAIL when it WAS and the file is
# missing. Skipping that case too would reproduce the same silent green under a wrong
# root -- the operator would have set SSBV2_ROOT, seen the suite pass, and still not
# have run these.
_needs_dataset = pytest.mark.skipif(
    not _SSBV2_ROOT, reason="SSBV2_ROOT is not set (see README, SSB v2 bootstrap)")


def _require_fixture():
    if not NO_CALCPR_INPUT.exists():
        pytest.fail(
            f"SSBV2_ROOT={_SSBV2_ROOT!r} is set but {NO_CALCPR_INPUT} does not exist. "
            f"These two tests are the only guard against _force_iterative_calc "
            f"regressing into a silent no-op; skipping here would hide a wrong root "
            f"behind a green suite.")

LOGGER = "src.agent_profiles.ssbv2_agent.scorer"


def _has_calcpr(p: Path) -> bool:
    return "<calcPr" in zipfile.ZipFile(p).read("xl/workbook.xml").decode("utf8", "replace")


@pytest.fixture
def warnings_captured():
    recs: list[str] = []

    class H(logging.Handler):
        def emit(self, r):      # noqa: D102
            recs.append(r.getMessage())

    h = H()
    lg = logging.getLogger(LOGGER)
    lg.addHandler(h)
    prev = lg.level
    lg.setLevel(logging.WARNING)
    yield recs
    lg.removeHandler(h)
    lg.setLevel(prev)


@_needs_dataset
def test_the_real_dataset_input_still_lacks_calcpr():
    """Guard the premise. If the dataset is ever reprocessed to include calcPr,
    the test below stops exercising the bug and would pass vacuously."""
    _require_fixture()
    assert not _has_calcpr(NO_CALCPR_INPUT), (
        "this input now HAS a <calcPr>; the None branch is no longer covered by "
        "this fixture — find another no-calcPr task or synthesise one explicitly")


@_needs_dataset
def test_missing_calcpr_gets_iteration_forced(tmp_path, warnings_captured):
    """The branch that used to raise. This is the whole point of the fix."""
    _require_fixture()
    p = tmp_path / "no_calcpr.xlsx"
    shutil.copy2(NO_CALCPR_INPUT, p)

    _force_iterative_calc(str(p))

    calc = openpyxl.load_workbook(p).calculation
    assert calc is not None, "calcPr was not created — the fix regressed"
    assert calc.iterate is True
    assert calc.iterateCount == 100
    assert calc.iterateDelta == 0.001
    assert not warnings_captured, (
        f"the fix path warned instead of working: {warnings_captured}")


def test_existing_iteration_settings_are_left_alone(tmp_path):
    """A workbook that ALREADY iterates must not be rewritten.

    The early return matters: golden files carry iterateDelta=0.0001 while this
    function's default is 0.001, so clobbering an explicit setting would change
    convergence on a file that was already correct.
    """
    p = tmp_path / "already.xlsx"
    wb = openpyxl.Workbook()
    wb.calculation = CalcProperties(calcId=1, iterate=True, iterateCount=7,
                                    iterateDelta=0.5)
    wb.save(p)

    _force_iterative_calc(str(p))

    calc = openpyxl.load_workbook(p).calculation
    assert (calc.iterateCount, calc.iterateDelta) == (7, 0.5)


def test_calcpr_present_but_iteration_off_is_turned_on(tmp_path):
    p = tmp_path / "off.xlsx"
    wb = openpyxl.Workbook()
    wb.calculation = CalcProperties(calcId=1, iterate=False)
    wb.save(p)

    _force_iterative_calc(str(p))

    calc = openpyxl.load_workbook(p).calculation
    assert calc.iterate is True
    assert calc.iterateCount == 100


def test_unreadable_file_warns_but_never_raises(tmp_path, warnings_captured):
    """Scoring must proceed. But the message has to name what was LOST.

    The old wording was "could not open/save", which is why 357 rollouts' worth
    of this warning class went unexamined in a log someone was reading.
    """
    p = tmp_path / "broken.xlsx"
    p.write_bytes(b"not a zip at all")

    _force_iterative_calc(str(p))          # must not raise

    assert warnings_captured, "a corrupt workbook produced no warning at all"
    msg = warnings_captured[-1]
    assert "NOT applied" in msg
    assert "circular-reference" in msg
