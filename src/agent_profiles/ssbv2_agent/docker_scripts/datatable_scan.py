"""Stage 1 (runs under the image's default python3.11, has openpyxl).

Scans a proc .xlsx for LibreOffice's post-recalc TABLE(...) formula strings
and emits a JSON list, one entry per cell, with everything needed to redo
the substitution.

Two representations exist for the same construct depending on when this
scan runs:
  - BEFORE any LibreOffice recalc: openpyxl exposes the anchor cell as a
    DataTableFormula OBJECT (ref/r1/r2 attrs); sibling cells are bare cached
    floats with no formula marker at all.
  - AFTER a LibreOffice --convert-to recalc round-trip: every cell in the
    range (anchor included) becomes the plain string
    '=TABLE($F$9;$C$12;$F10;$C$15;G$9)' -- confirmed empirically. This is
    fully self-describing PER CELL (corner, row-input, row-header,
    col-input, col-header all as literal cell refs in the string), so this
    scan targets that string form directly rather than needing the pre-recalc
    file -- the scorer pipeline always calls this AFTER
    _recalculate_with_libreoffice, so the string form is what is actually on
    disk when this runs.

Usage: python3 datatable_scan.py <xlsx_path> <output_json_path>
"""
import json
import re
import sys

import openpyxl

# Separator is locale-dependent (comma or semicolon observed); tolerate either.
TABLE_RE = re.compile(
    r"^=TABLE\(\s*([^,;]+)\s*[,;]\s*([^,;]+)\s*[,;]\s*([^,;]+)\s*[,;]\s*([^,;]+)\s*[,;]\s*([^,;)]+)\s*\)$",
    re.IGNORECASE,
)


def scan(xlsx_path: str) -> list[dict]:
    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    cells = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not isinstance(v, str):
                    continue
                m = TABLE_RE.match(v.strip())
                if not m:
                    continue
                corner, row_input, row_header, col_input, col_header = (
                    g.strip().lstrip("$").replace("$", "") for g in m.groups()
                )
                cells.append({
                    "sheet": sheet_name,
                    "cell": cell.coordinate,
                    "corner": corner,
                    "row_input": row_input,
                    "row_header": row_header,
                    "col_input": col_input,
                    "col_header": col_header,
                })
    return cells


if __name__ == "__main__":
    xlsx_path, out_json = sys.argv[1], sys.argv[2]
    cells = scan(xlsx_path)
    with open(out_json, "w") as f:
        json.dump(cells, f)
    print(f"Found {len(cells)} TABLE() cell(s)")
